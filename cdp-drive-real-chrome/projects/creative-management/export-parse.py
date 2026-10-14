"""AC-17 导出文件逐行解析（去标识化）。

用法：python -I export-parse.py <xlsx> <输出前缀>
产出：
  <前缀>-rows.csv     每个数据行一行：行号、Ad Type、Ad 列是否含 `ID:`、是否等于 `Creative`、Ad 列文字哈希（sha1 前 10 位）、
                      Ad 列行数、各列锚定图片数（img_c0 / img_c1 …，Excel 0 基列号）
  <前缀>-summary.json 按 Ad Type 汇总：行数、含 ID 行数、等于 Creative 行数、每行图片数分布；全表图片总数；
                      「SP 行应有 2 张（缩略图 + 商品图）」的差额逐行列出（只给行号与哈希）
不落任何原文（Ad 名称、商品 Title、ASIN 都只留哈希），原文件由调用方验完删除。
"""
import collections, csv, hashlib, json, sys
from openpyxl import load_workbook

src, prefix = sys.argv[1], sys.argv[2]
wb = load_workbook(src)
ws = wb.worksheets[0]

hdr_row, cols = None, {}
for r in range(1, 6):
    vals = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
    if "Ad Type" in vals:
        hdr_row, cols = r, {v: i for i, v in enumerate(vals) if v}
        break
if hdr_row is None:
    sys.exit("找不到含 Ad Type 的表头行")
ad_col = 0  # 首列为 Ad（两条导出链路一致）
type_col = cols["Ad Type"]

imgs = collections.defaultdict(lambda: collections.Counter())  # excel 行号(1 基) -> {列号: 张数}
total_imgs = 0
for im in ws._images:
    a = im.anchor._from
    imgs[a.row + 1][a.col] += 1
    total_imgs += 1
max_img_col = max([c for cnt in imgs.values() for c in cnt] or [0])

rows = []
for r in range(hdr_row + 1, ws.max_row + 1):
    ad = ws.cell(r, ad_col + 1).value
    typ = ws.cell(r, type_col + 1).value
    if ad is None and typ is None:
        continue
    text = "" if ad is None else str(ad)
    row = {"excel_row": r, "ad_type": "" if typ is None else str(typ).strip(),
           "has_id_line": any(line.strip().startswith("ID:") for line in text.split("\n")),
           "equals_creative": text.strip() == "Creative",
           "ad_text_sha1": hashlib.sha1(text.encode("utf-8")).hexdigest()[:10],
           "ad_lines": len(text.split("\n")) if text else 0}
    for c in range(max_img_col + 1):
        row[f"img_c{c}"] = imgs.get(r, {}).get(c, 0)
    rows.append(row)

with open(prefix + "-rows.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
    w.writeheader(); w.writerows(rows)

by = collections.defaultdict(lambda: {"rows": 0, "has_id_line": 0, "equals_creative": 0, "imgs_per_row": collections.Counter()})
for x in rows:
    b = by[x["ad_type"]]
    b["rows"] += 1; b["has_id_line"] += x["has_id_line"]; b["equals_creative"] += x["equals_creative"]
    b["imgs_per_row"][sum(x[k] for k in x if k.startswith("img_c"))] += 1
img_rows_outside = sorted(set(imgs) - {x["excel_row"] for x in rows})
sp = [x for x in rows if x["ad_type"] == "SP"]
sp_short = [{"excel_row": x["excel_row"], "imgs": sum(x[k] for k in x if k.startswith("img_c")), "has_id_line": x["has_id_line"], "ad_text_sha1": x["ad_text_sha1"]}
            for x in sp if sum(x[k] for k in x if k.startswith("img_c")) != 2]
summary = {"file_sheet": ws.title, "header_row": hdr_row, "header": list(cols)[:8], "data_rows": len(rows), "total_images": total_imgs,
           "image_columns": sorted({c for cnt in imgs.values() for c in cnt}),
           "images_on_non_data_rows": img_rows_outside[:20],
           "by_ad_type": {k: {**v, "imgs_per_row": dict(sorted(v["imgs_per_row"].items()))} for k, v in sorted(by.items())},
           "sp_rows_without_exactly_2_images": {"count": len(sp_short), "rows": sp_short[:200]},
           "sp_rows_without_id_line": [x["excel_row"] for x in sp if not x["has_id_line"]][:50]}
with open(prefix + "-summary.json", "w", encoding="utf-8") as f:
    json.dump(summary, f, ensure_ascii=False, indent=1)
print(json.dumps({k: summary[k] for k in ("data_rows", "total_images", "image_columns", "by_ad_type")}, ensure_ascii=False))
print("sp_rows_without_exactly_2_images:", len(sp_short))
